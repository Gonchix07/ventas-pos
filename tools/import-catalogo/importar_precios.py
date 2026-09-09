"""
Carga precios.xlsx en las listas de precios de POS-Ventas.

El archivo trae UNA fila por artículo y una columna por lista (azul / roja / folder), con el
precio de la UNIDAD SUELTA. Cada presentación se valoriza multiplicando por sus unidades por
bulto (de `Presentaciones.UnidadXBulto`, la fuente de verdad del catálogo) — la misma regla que
usa el editor de precios de la app (Pos.Domain.Services.PrecioPorBulto).

Cómo se determinó que el precio es unitario y no del bulto: el mismo aceite AC NATURA aparece en
4 packs distintos (15x0.9L, 12x1.5L, 6x3L, 4x5L) y solo bajo la hipótesis "por unidad" el precio
por litro converge (~$3.800–4.000/L); bajo "por bulto" da entre $267 y $950/L.

Un 0 o una celda vacía significan "sin precio" y no se cargan (decisión del usuario, 2026-09-08:
tratar blanco igual que 0).

La columna `unidxbult` del archivo es la unidad por bulto SEGÚN EL ERP viejo — no se usa para
calcular (se sigue usando la de `Presentaciones`, que es la que ve la app), pero se compara y se
avisa si difiere, por si el ERP y el catálogo real quedaron desalineados.

Uso:  python importar_precios.py [--dry-run]
"""
import sys
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
import pyodbc

ORIGEN = r"D:\Escritorio\precios.xlsx"

CONN = ("DRIVER={ODBC Driver 17 for SQL Server};SERVER=192.168.4.9;DATABASE=POS-Ventas;"
        "UID=claude;PWD=3edcHG9ijn;TrustServerCertificate=yes;Encrypt=yes")

# columna del archivo -> CodigoInterno de la lista en la BD
LISTAS = {"azul": "AZUL", "roja": "ROJA", "folder": "FOLDER AGO"}

AHORA = datetime.now(timezone.utc).replace(tzinfo=None)
POR = "import:precios.xlsx"


def redondear(valor):
    """2 decimales, medio hacia arriba: igual que PrecioPorBulto en el dominio."""
    return Decimal(str(valor)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def leer_precios():
    wb = openpyxl.load_workbook(ORIGEN, data_only=True, read_only=True)
    ws = wb["precios"]
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    cols = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert cols == ["codigo", "descrip", "unidxbult", "azul", "roja", "folder"], f"Cabecera inesperada: {cols}"

    por_codigo: dict[str, dict] = {}
    for codigo, descrip, unidxbult, azul, roja, folder in filas:
        cod = str(codigo).strip().lstrip("0") or "0"
        por_codigo[cod] = {  # el último gana si el código está repetido en el archivo
            "unidxbult": unidxbult,
            "azul": azul or 0,
            "roja": roja or 0,
            "folder": folder or 0,
        }
    return por_codigo


def main(dry_run=False):
    print("Leyendo precios.xlsx…")
    precios = leer_precios()
    print(f"  {len(precios):,} artículos con fila de precio")

    cx = pyodbc.connect(CONN, autocommit=False)
    cur = cx.cursor()
    try:
        cur.execute("SELECT CodigoInterno, IdListaPrecio FROM ListasPrecios")
        id_lista = {c: i for c, i in cur.fetchall()}
        faltan = [n for n in LISTAS.values() if n not in id_lista]
        if faltan:
            raise SystemExit(f"No existen estas listas en la BD: {faltan}")

        cur.execute("""SELECT a.CodigoInterno, a.IdArticulo, p.IdPresentacion, p.UnidadXBulto
                       FROM Articulos a JOIN Presentaciones p ON p.IdArticulo = a.IdArticulo""")
        pres = {}
        for cod, id_art, id_pres, uxb in cur.fetchall():
            pres.setdefault(cod, []).append((id_art, id_pres, float(uxb)))
        print(f"  {len(pres):,} artículos en el catálogo con presentaciones")

        sin_articulo = {c: v for c, v in precios.items() if c not in pres}
        con_precio = {c: v for c, v in sin_articulo.items() if v["azul"] or v["roja"] or v["folder"]}
        print(f"  {len(sin_articulo):,} filas sin artículo en el catálogo "
              f"({len(con_precio):,} de ellas traían precio y se pierden)")

        desalineados = 0
        for cod, v in precios.items():
            if cod in pres and v["unidxbult"]:
                unidades_reales = {uxb for _, _, uxb in pres[cod]}
                if float(v["unidxbult"]) not in unidades_reales:
                    desalineados += 1
        if desalineados:
            print(f"  ! {desalineados:,} artículos donde `unidxbult` del archivo no coincide con "
                  f"ninguna presentación real del catálogo (se usa igual la del catálogo)")

        filas = []
        por_lista = {}
        for col, nombre in LISTAS.items():
            id_l = id_lista[nombre]
            n = 0
            for cod, v in precios.items():
                unitario = v[col]
                if not unitario or unitario <= 0:  # 0 o blanco = sin precio
                    continue
                for id_art, id_pres, uxb in pres.get(cod, ()):
                    filas.append((id_l, id_pres, id_art, float(redondear(unitario * uxb)), 0.0, AHORA, POR))
                    n += 1
            por_lista[nombre] = n
            print(f"  {nombre}: {n:,} precios de presentación")

        print(f"\nTotal a insertar: {len(filas):,}")
        if dry_run:
            print("--dry-run: no se toca la base.")
            return

        cur.execute("DELETE FROM Precios")
        print(f"Precios previos borrados: {cur.rowcount}")

        cur.fast_executemany = True
        sql = ("INSERT INTO Precios (IdListaPrecio, IdPresentacion, IdArticulo, PrecioFinal, "
               "ImpuestoInterno, CreatedAtUtc, CreatedBy) VALUES (?,?,?,?,?,?,?)")
        for i in range(0, len(filas), 2000):
            cur.executemany(sql, filas[i:i + 2000])

        cx.commit()
        print("\nOK — commit hecho.")
    except Exception:
        cx.rollback()
        print("\nERROR — rollback, la base quedó como estaba.")
        raise
    finally:
        cx.close()


if __name__ == "__main__":
    main(dry_run="--dry-run" in sys.argv)
