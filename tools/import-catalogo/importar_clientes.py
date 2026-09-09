"""
Actualiza Clientes desde D:\\Escritorio\\Clientes.xlsx (export completo del ERP): actualiza los
que ya existen (por CodigoInt) e inserta los que son nuevos.

Formato del archivo (una fila por cliente, o varias si tiene más de una tarjeta a lo largo del
tiempo — ver más abajo): `codigo_a, nombre_a, domicilio_fis, copostal, localidad, email, cuit,
numdoc, nombre_b, presupuesto, inactivo, tipo_tarjeta, codigo_b`.

Decisiones tomadas con el usuario (2026-09-08):
  - `nombre_b` (Condición de IVA, texto) se mapea por descripción a `CondicionesIva` — mapeo
    verificado cruzando clientes ya cargados contra su IdCondIva actual en la BD.
  - `inactivo` viene INVERTIDO respecto a `Cliente.Activo` (1 = inactivo).
  - `tipo_tarjeta` viene como código del ERP ('03'/'04'), no como texto — verificado cruzando
    contra `TarjetasClientes` ya cargadas: 03 = TARJETA ROJA (IdTipoTarjeta 4), 04 = TARJETA AZUL (5).
  - **Códigos de cliente repetidos en el archivo** (2.844 casos, 2-3 filas): el archivo trae
    historial de tarjetas, no una fila por cliente. Se toma la ÚLTIMA fila de cada código como la
    vigente (se asume orden cronológico) y esa es la única tarjeta que se deja Activa=True — se
    anula (Activa=False, FechaBajaUtc=ahora) cualquier otra tarjeta que el cliente tuviera activa
    en la BD. Esto de paso corrige el problema previo de clientes con 2+ tarjetas activas a la vez
    (2.310 casos, arrastrados de la carga original "importacion-padron").
  - `email`/`cuit` sucios (placeholders del ERP: "@", ".", "-", "no tiene"...) NO pisan un valor
    bueno que ya estuviera cargado: si el valor del archivo no parece válido, se deja como estaba
    (mismo criterio que `importar_fantasia.py`). Los demás campos (nombre, domicilio, localidad,
    cod. postal, condición IVA, presupuesto, activo) SÍ se actualizan siempre con lo del archivo,
    porque son campos limpios del ERP.
  - `numdoc` puede venir como fórmula de Excel (`=MID(...)`) en algunas filas de tabla — hay que
    leer con `data_only=True` para traer el valor ya calculado, no la fórmula en texto.

Uso:  python importar_clientes.py [--dry-run]
"""
import re
import sys
from datetime import datetime, timezone

import openpyxl
import pyodbc

ORIGEN = r"D:\Escritorio\Clientes.xlsx"
CONN = ("DRIVER={ODBC Driver 17 for SQL Server};SERVER=192.168.4.9;DATABASE=POS-Ventas;"
        "UID=claude;PWD=3edcHG9ijn;TrustServerCertificate=yes;Encrypt=yes")

AHORA = datetime.now(timezone.utc).replace(tzinfo=None)
POR = "import:Clientes.xlsx"

CONDICION_IVA = {
    "CONSUMIDOR FINAL": 4,
    "REGIMEN SIMPLIFICADO": 2,
    "RESPONSABLE INSCRIPTO": 1,
    "RESP. INS M": 1,
    "EXENTO/NO ALCANZADO": 3,
    "RESPONSABLE NO INSCRIPTO": 9,
    "SUJETO NO CATEGORIZADO": 10,
}
TIPO_TARJETA = {"03": 4, "04": 5}  # ROJA / AZUL

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def limpiar_cuit(valor) -> str | None:
    if not valor:
        return None
    digitos = re.sub(r"\D", "", str(valor))
    return digitos if len(digitos) == 11 else None


def limpiar_email(valor) -> str | None:
    if not valor:
        return None
    v = str(valor).strip()
    return v[:120] if EMAIL_RE.match(v) else None


class Fila:
    __slots__ = ("codigo", "nombre", "domicilio", "cp", "localidad", "email", "cuit", "documento",
                 "id_condiva", "presupuesto", "activo", "id_tipo_tarjeta", "nro_tarjeta")


def leer():
    wb = openpyxl.load_workbook(ORIGEN, data_only=True, read_only=True)
    ws = wb["clientes"]
    filas_por_codigo: dict[str, Fila] = {}
    sin_condiva, sin_tarjeta_valida = set(), 0
    total = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        total += 1
        cod = str(r[0]).strip()
        f = Fila()
        f.codigo = cod
        f.nombre = (str(r[1]).strip() if r[1] else "") or cod
        f.domicilio = str(r[2]).strip() if r[2] else None
        f.cp = str(r[3]).strip() if r[3] else None
        f.localidad = str(r[4]).strip() if r[4] else None
        f.email = limpiar_email(r[5])
        f.cuit = limpiar_cuit(r[6])
        f.documento = str(r[7]).strip() if r[7] else None
        nombre_b = (str(r[8]).strip().upper() if r[8] else "")
        f.id_condiva = CONDICION_IVA.get(nombre_b)
        if r[8] and f.id_condiva is None:
            sin_condiva.add(nombre_b)
        f.presupuesto = str(r[9]).strip() == "1"
        f.activo = str(r[10]).strip() != "1"  # "inactivo" viene invertido
        tt = str(r[11]).strip() if r[11] else ""
        cb = str(r[12]).strip() if r[12] else ""
        if tt in TIPO_TARJETA and cb:
            f.id_tipo_tarjeta = TIPO_TARJETA[tt]
            f.nro_tarjeta = cb
        else:
            f.id_tipo_tarjeta = None
            f.nro_tarjeta = None
            if tt or cb:
                sin_tarjeta_valida += 1
        # clave insensible a mayúsculas: el índice único de CodigoInt en la BD lo es (collation
        # por defecto de SQL Server), y el archivo trae el mismo código en distinta caja alguna vez
        # (ej. "c9840" en el archivo vs "C9840" ya cargado) — sin esto se intentaba insertar como
        # nuevo un cliente que ya existía, y violaba el índice único.
        filas_por_codigo[cod.upper()] = f  # última fila del código gana (se asume orden cronológico)
    return total, filas_por_codigo, sin_condiva, sin_tarjeta_valida


def main(dry_run=False):
    print("Leyendo Clientes.xlsx…")
    total, filas, sin_condiva, sin_tarjeta_valida = leer()
    print(f"  {total:,} filas → {len(filas):,} clientes únicos (última fila de cada código)")
    if sin_condiva:
        print(f"  ! condición de IVA sin mapeo (se deja NULL, no debería pasar): {sin_condiva}")
    if sin_tarjeta_valida:
        print(f"  {sin_tarjeta_valida:,} filas con tipo/código de tarjeta incompleto (se ignora la tarjeta)")

    cx = pyodbc.connect(CONN, autocommit=False)
    cur = cx.cursor()
    try:
        cur.execute("SELECT IdCliente, CodigoInt, Descripcion, Domicilio, CodigoPostal, Localidad, "
                     "Email, Cuit, Documento, IdCondIva, PermitePresupuesto, Activo FROM Clientes")
        actuales = {r.CodigoInt.strip().upper(): r for r in cur.fetchall()}
        print(f"Clientes en la BD: {len(actuales):,}")

        cur.execute("SELECT IdCliente, IdTipoTarjeta, NroTarjeta FROM TarjetasClientes WHERE Activa=1")
        tarjeta_activa_por_cliente: dict[int, list] = {}
        for id_cli, id_tt, nro in cur.fetchall():
            tarjeta_activa_por_cliente.setdefault(id_cli, []).append((id_tt, nro))

        nuevos, updates_cliente = [], []
        anular_tarjetas, insertar_tarjetas, tarjetas_sin_cambio = [], [], 0
        sin_condiva_final = 0

        for cod, f in filas.items():
            existente = actuales.get(cod)
            if existente is None:
                nuevos.append(f)
                continue

            if f.id_condiva is None:
                f.id_condiva = existente.IdCondIva
                sin_condiva_final += 1
            email_final = f.email if f.email is not None else existente.Email
            cuit_final = f.cuit if f.cuit is not None else existente.Cuit

            updates_cliente.append((
                f.nombre, f.domicilio, f.cp, f.localidad, email_final, cuit_final, f.documento,
                f.id_condiva, f.presupuesto, f.activo, AHORA, POR, existente.IdCliente,
            ))

            id_cliente = existente.IdCliente
            activas = tarjeta_activa_por_cliente.get(id_cliente, [])
            if f.id_tipo_tarjeta is None:
                continue  # el archivo no trae tarjeta para este cliente: no se toca lo que hay
            ya_es_la_vigente = any(t == (f.id_tipo_tarjeta, f.nro_tarjeta) for t in activas)
            if ya_es_la_vigente and len(activas) == 1:
                tarjetas_sin_cambio += 1
                continue
            # anular todas las activas que no sean exactamente la que queda vigente
            for id_tt, nro in activas:
                if (id_tt, nro) != (f.id_tipo_tarjeta, f.nro_tarjeta):
                    anular_tarjetas.append((AHORA, AHORA, POR, id_cliente, id_tt, nro))
            if not ya_es_la_vigente:
                insertar_tarjetas.append((id_cliente, f.id_tipo_tarjeta, f.nro_tarjeta, AHORA, POR))

        tarjetas_de_nuevos = sum(1 for f in nuevos if f.id_tipo_tarjeta is not None)
        print(f"\nClientes nuevos a insertar: {len(nuevos):,} ({tarjetas_de_nuevos:,} con tarjeta)")
        print(f"Clientes existentes a actualizar: {len(updates_cliente):,} "
              f"(condición IVA no mapeada, se mantuvo la que tenía: {sin_condiva_final:,})")
        print(f"Tarjetas de clientes existentes: sin cambio {tarjetas_sin_cambio:,} | "
              f"a anular {len(anular_tarjetas):,} | a insertar como nueva vigente {len(insertar_tarjetas):,}")

        if dry_run:
            print("\n--dry-run: no se toca la base.")
            return

        cur.fast_executemany = True

        if nuevos:
            sql = ("INSERT INTO Clientes (CodigoInt, Descripcion, Domicilio, CodigoPostal, Localidad, "
                   "Email, Cuit, Documento, IdCondIva, PermitePresupuesto, Activo, "
                   "AdmiteCuentaCorriente, CreatedAtUtc, CreatedBy) "
                   "VALUES (?,?,?,?,?,?,?,?,?,?,?,0,?,?)")
            filas_ins = [(f.codigo, f.nombre, f.domicilio, f.cp, f.localidad, f.email, f.cuit,
                          f.documento, f.id_condiva, f.presupuesto, f.activo, AHORA, POR)
                         for f in nuevos]
            for i in range(0, len(filas_ins), 1000):
                cur.executemany(sql, filas_ins[i:i + 1000])

        if updates_cliente:
            sql = ("UPDATE Clientes SET Descripcion=?, Domicilio=?, CodigoPostal=?, Localidad=?, "
                   "Email=?, Cuit=?, Documento=?, IdCondIva=?, PermitePresupuesto=?, Activo=?, "
                   "UpdatedAtUtc=?, UpdatedBy=? WHERE IdCliente=?")
            for i in range(0, len(updates_cliente), 1000):
                cur.executemany(sql, updates_cliente[i:i + 1000])

        if anular_tarjetas:
            sql = ("UPDATE TarjetasClientes SET Activa=0, FechaBajaUtc=?, UpdatedAtUtc=?, UpdatedBy=? "
                   "WHERE IdCliente=? AND IdTipoTarjeta=? AND NroTarjeta=? AND Activa=1")
            for i in range(0, len(anular_tarjetas), 1000):
                cur.executemany(sql, anular_tarjetas[i:i + 1000])

        # las tarjetas de los clientes NUEVOS se insertan recién acá, porque hasta ahora no
        # tenían IdCliente (identity) — se resuelve con el mismo cruce por CodigoInt.
        if nuevos:
            cur.execute("SELECT IdCliente, CodigoInt FROM Clientes")
            id_por_codigo = {c.strip(): i for i, c in cur.fetchall()}
            for f in nuevos:
                if f.id_tipo_tarjeta is not None:
                    insertar_tarjetas.append((id_por_codigo[f.codigo], f.id_tipo_tarjeta, f.nro_tarjeta, AHORA, POR))

        if insertar_tarjetas:
            sql = ("INSERT INTO TarjetasClientes (IdCliente, IdTipoTarjeta, NroTarjeta, Activa, "
                   "CreatedAtUtc, CreatedBy) VALUES (?,?,?,1,?,?)")
            for i in range(0, len(insertar_tarjetas), 1000):
                cur.executemany(sql, insertar_tarjetas[i:i + 1000])

        cx.commit()
        print("\nOK — commit hecho.")
        cur.execute("SELECT COUNT(*) FROM Clientes")
        print(f"Total clientes en la BD ahora: {cur.fetchone()[0]:,}")
    except Exception:
        cx.rollback()
        print("\nERROR — rollback, la base quedó como estaba.")
        raise
    finally:
        cx.close()


if __name__ == "__main__":
    main(dry_run="--dry-run" in sys.argv)
